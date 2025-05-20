from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    UnstructuredPowerPointLoader,
    TextLoader
)
from langchain.text_splitter import RecursiveCharacterTextSplitter
import tempfile
import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from langchain.schema import Document

class DocumentService:
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=2000,  # Memperbesar chunk size
            chunk_overlap=200,
            length_function=len,
            is_separator_regex=False,
            add_start_index=True  # Menambahkan index untuk tracking
        )
    
    def process_file(self, file):
        if not file or not file.filename:
            print("File tidak valid")
            return None
        
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, file.filename)
        
        try:
            # Simpan file
            file.save(temp_path)
            
            print(f"Processing file: {file.filename}")
            
            # Khusus untuk PowerPoint, coba ekstrak langsung
            if file.filename.lower().endswith('.pptx'):
                try:
                    prs = Presentation(temp_path)
                    all_content = []
                    
                    for slide_number, slide in enumerate(prs.slides, 1):
                        slide_content = []
                        # Tambahkan nomor slide
                        slide_content.append(f"Slide {slide_number}:")
                        
                        # Ekstrak teks dari shape
                        for shape in slide.shapes:
                            # Teks dari shape
                            if hasattr(shape, "text") and shape.text.strip():
                                slide_content.append(shape.text.strip())
                            
                            # Teks dari tabel
                            if shape.has_table:
                                table_text = []
                                for row in shape.table.rows:
                                    row_text = []
                                    for cell in row.cells:
                                        if cell.text.strip():
                                            row_text.append(cell.text.strip())
                                    if row_text:
                                        table_text.append(" | ".join(row_text))
                                if table_text:
                                    slide_content.append("\n".join(table_text))
                            
                            # Deskripsi untuk gambar
                            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                                slide_content.append("[Gambar terdeteksi]")
                        
                        if slide_content:
                            all_content.append("\n".join(slide_content))
                    
                    if all_content:
                        doc = Document(
                            page_content="\n\n".join(all_content),
                            metadata={"source": file.filename}
                        )
                        texts = self.text_splitter.split_documents([doc])
                        print(f"Berhasil memproses PowerPoint: {len(texts)} bagian")
                        return texts
                    else:
                        print("Presentasi hanya berisi gambar tanpa teks")
                        # Buat dokumen dengan deskripsi gambar
                        doc = Document(
                            page_content="Presentasi ini berisi terutama gambar dan konten visual",
                            metadata={"source": file.filename}
                        )
                        return [doc]
                        
                except Exception as e:
                    print(f"Error saat memproses PowerPoint: {str(e)}")
            
            # Untuk file lain, gunakan loader standar
            loader = self._get_loader(temp_path, file.filename)
            if not loader:
                return None
                
            documents = loader.load()
            if not documents:
                return None
                
            texts = self.text_splitter.split_documents(documents)
            if not texts:
                print("Tidak ada teks yang bisa diekstrak")
                return None
                
            print(f"Berhasil memproses dokumen: {len(texts)} bagian")
            return texts
            
        except Exception as e:
            print(f"Error saat memproses file: {str(e)}")
            return None
            
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            if os.path.exists(temp_dir):
                os.rmdir(temp_dir)
    
    def _get_loader(self, file_path, filename):
        if not os.path.exists(file_path):
            print(f"File tidak ditemukan: {file_path}")
            return None
            
        try:
            ext = filename.lower()
            if ext.endswith('.pdf'):
                return PyPDFLoader(file_path)
            elif ext.endswith('.docx'):
                return Docx2txtLoader(file_path)
            elif ext.endswith('.pptx'):
                return UnstructuredPowerPointLoader(file_path, mode="elements")
            elif ext.endswith('.txt'):
                return TextLoader(file_path, encoding='utf-8')
            else:
                print(f"Format file tidak didukung: {filename}")
                return None
        except Exception as e:
            print(f"Error saat membuat loader: {str(e)}")
            return None